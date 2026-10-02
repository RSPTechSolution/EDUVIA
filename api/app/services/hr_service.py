import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import uuid
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from app.models.hr import SalarySlip, Staff

async def export_bank_advice_excel(db: AsyncSession, payroll_run_id: uuid.UUID) -> bytes:
    # Get slips
    slips = await db.scalars(
        select(SalarySlip)
        .options(selectinload(SalarySlip.staff))
        .where(SalarySlip.payroll_run_id == payroll_run_id)
    )
    slips = list(slips.all())
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bank Transfer Advice"

    # Header row
    headers = ["S.No", "Staff Name", "Employee ID", "Bank", "Account No", "IFSC Code", "Net Salary"]
    header_fill = PatternFill("solid", fgColor="1A3C5E")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for i, slip in enumerate(slips, 2):
        ws.append([
            i - 1, f"{slip.staff.first_name} {slip.staff.last_name}",
            slip.staff.employee_id, slip.staff.bank_name,
            slip.staff.account_no, slip.staff.ifsc_code, float(slip.net_salary),
        ])

    # Auto-width
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value or "")) for c in col) + 4

    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
